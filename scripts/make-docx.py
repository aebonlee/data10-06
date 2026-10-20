"""기획서 Markdown → Word(docx): python3 scripts/make-docx.py
docs/01_프로젝트_기획서.md 를 읽어 docs/01_프로젝트_기획서.docx 로 씁니다(python-docx 필요).
제목·표·글머리·번호·인용·코드(mermaid 그림 정의는 글자 그대로)·**굵게** 만 다룹니다."""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
src = root / 'docs' / '01_프로젝트_기획서.md'
out = root / 'docs' / '01_프로젝트_기획서.docx'
FONT = '맑은 고딕'

doc = Document()
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
for name in ['Heading 1', 'Heading 2', 'Heading 3', 'Title']:
    s = doc.styles[name]; s.font.name = FONT; s.element.rPr.rFonts.set(qn('w:eastAsia'), FONT); s.font.color.rgb = RGBColor(0x0F, 0x25, 0x44)

def runs(p, text):
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not part: continue
        r = p.add_run(part.replace('`', '')); r.bold = i % 2 == 1
        r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

lines = src.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith('```'):
        i += 1; code = []
        while i < len(lines) and not lines[i].startswith('```'):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph()
        r = p.add_run('\n'.join(code)); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
        i += 1; continue
    if ln.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?', c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = ''
                runs(cell.paragraphs[0], val)
                if r == 0:
                    for rr in cell.paragraphs[0].runs: rr.bold = True
        doc.add_paragraph(); continue
    m = re.match(r'^(#{1,3})\s+(.*)', ln)
    if m:
        lvl = len(m.group(1)); doc.add_heading(m.group(2), level=0 if lvl == 1 else lvl - 1)
    elif ln.strip() == '---':
        pass
    elif re.match(r'^\s*[-*]\s+', ln):
        runs(doc.add_paragraph(style='List Bullet'), re.sub(r'^\s*[-*]\s+', '', ln))
    elif re.match(r'^\d+\.\s+', ln):
        runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\.\s+', '', ln))
    elif ln.startswith('>'):
        p = doc.add_paragraph(); runs(p, ln.lstrip('> ').strip())
        for r in p.runs: r.italic = True
    elif ln.strip():
        runs(doc.add_paragraph(), ln)
    i += 1
doc.save(out)
print(out.relative_to(root))
